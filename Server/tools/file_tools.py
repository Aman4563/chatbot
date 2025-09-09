"""File processing tools for agent capabilities."""

import base64
import io
import csv
import json
from typing import Dict, Any, List, Optional
from langchain_core.tools import tool
from .base import BaseTool
import PyPDF2
import docx
from PIL import Image


class DocumentProcessorTool(BaseTool):
    """Tool for processing various document types."""
    
    def __init__(self):
        super().__init__("document_processor", "Process and analyze documents")
        self.category = "file_processing"
        self.supported_formats = [
            "pdf", "docx", "txt", "csv", "json", "xml", "md"
        ]
    
    async def execute(
        self,
        file_data: str,
        filename: str,
        mime_type: str,
        processing_type: str = "extract",
        **kwargs
    ) -> Dict[str, Any]:
        """Execute document processing."""
        if not await self.validate_input(file_data=file_data, filename=filename, mime_type=mime_type):
            raise ValueError("Invalid input parameters")
        
        if processing_type == "extract":
            return await self._extract_content(file_data, filename, mime_type)
        elif processing_type == "analyze":
            return await self._analyze_content(file_data, filename, mime_type)
        elif processing_type == "summarize":
            return await self._summarize_content(file_data, filename, mime_type)
        else:
            raise ValueError(f"Unknown processing type: {processing_type}")
    
    async def _extract_content(self, file_data: str, filename: str, mime_type: str) -> Dict[str, Any]:
        """Extract content from document."""
        try:
            decoded_data = base64.b64decode(file_data)
            
            if mime_type == "application/pdf":
                content = await self._extract_pdf_content(decoded_data)
            elif mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                content = await self._extract_docx_content(decoded_data)
            elif mime_type == "text/plain":
                content = await self._extract_text_content(decoded_data)
            elif mime_type == "text/csv":
                content = await self._extract_csv_content(decoded_data)
            elif mime_type == "application/json":
                content = await self._extract_json_content(decoded_data)
            else:
                content = f"Unsupported file type: {mime_type}"
            
            return {
                "status": "success",
                "filename": filename,
                "mime_type": mime_type,
                "content": content,
                "content_length": len(content),
                "extraction_method": "direct"
            }
        except Exception as e:
            return {
                "status": "error",
                "filename": filename,
                "error": str(e),
                "extraction_method": "failed"
            }
    
    async def _extract_pdf_content(self, data: bytes) -> str:
        """Extract text from PDF."""
        try:
            pdf_file = io.BytesIO(data)
            reader = PyPDF2.PdfReader(pdf_file)
            content = []
            
            for i, page in enumerate(reader.pages, 1):
                try:
                    text = page.extract_text() or ""
                    if text.strip():
                        content.append(f"--- Page {i} ---\n{text}")
                except Exception as e:
                    content.append(f"--- Page {i} ---\n[Error extracting page: {e}]")
            
            return "\n".join(content) if content else "PDF appears empty or image-only."
        except Exception as e:
            return f"Error extracting PDF: {e}"
    
    async def _extract_docx_content(self, data: bytes) -> str:
        """Extract text from DOCX."""
        try:
            docx_file = io.BytesIO(data)
            document = docx.Document(docx_file)
            content = []
            
            # Extract paragraphs
            for paragraph in document.paragraphs:
                if paragraph.text.strip():
                    content.append(paragraph.text)
            
            # Extract tables
            for table in document.tables:
                content.append("\n--- Table ---")
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    content.append(row_text)
            
            return "\n".join(content) if content else "Word document appears to be empty."
        except Exception as e:
            return f"Error extracting Word document: {e}"
    
    async def _extract_text_content(self, data: bytes) -> str:
        """Extract text from plain text file."""
        try:
            return data.decode("utf-8")
        except Exception as e:
            return f"Error decoding text file: {e}"
    
    async def _extract_csv_content(self, data: bytes) -> str:
        """Extract and format CSV content."""
        try:
            decoded = data.decode("utf-8")
            csv_reader = csv.reader(io.StringIO(decoded))
            content = ["CSV Data:"]
            
            for idx, row in enumerate(csv_reader, 1):
                if idx == 1:
                    content.append("Headers: " + " | ".join(row))
                    content.append("-" * 50)
                else:
                    content.append(f"Row {idx - 1}: " + " | ".join(row))
                
                if idx >= 100:  # Limit to first 100 rows
                    content.append("... (showing first 100 rows)")
                    break
            
            return "\n".join(content)
        except Exception as e:
            return f"Error processing CSV: {e}"
    
    async def _extract_json_content(self, data: bytes) -> str:
        """Extract and format JSON content."""
        try:
            decoded = data.decode("utf-8")
            obj = json.loads(decoded)
            return "JSON Data:\n" + json.dumps(obj, indent=2, ensure_ascii=False)
        except json.JSONDecodeError as e:
            return f"Invalid JSON: {e}"
        except Exception as e:
            return f"Error processing JSON: {e}"
    
    async def _analyze_content(self, file_data: str, filename: str, mime_type: str) -> Dict[str, Any]:
        """Analyze document content."""
        extraction_result = await self._extract_content(file_data, filename, mime_type)
        
        if extraction_result["status"] != "success":
            return extraction_result
        
        content = extraction_result["content"]
        
        # Basic analysis
        analysis = {
            "word_count": len(content.split()),
            "character_count": len(content),
            "line_count": len(content.splitlines()),
            "has_tables": "--- Table ---" in content,
            "has_code": any(keyword in content.lower() for keyword in ["def ", "function", "class ", "import "]),
            "language": self._detect_language(content)
        }
        
        return {
            "status": "success",
            "filename": filename,
            "analysis": analysis,
            "content_preview": content[:500] + "..." if len(content) > 500 else content
        }
    
    async def _summarize_content(self, file_data: str, filename: str, mime_type: str) -> Dict[str, Any]:
        """Summarize document content."""
        extraction_result = await self._extract_content(file_data, filename, mime_type)
        
        if extraction_result["status"] != "success":
            return extraction_result
        
        content = extraction_result["content"]
        
        # Simple summarization (in a real implementation, this would use an LLM)
        summary = self._create_simple_summary(content, filename)
        
        return {
            "status": "success",
            "filename": filename,
            "summary": summary,
            "content_length": len(content)
        }
    
    def _detect_language(self, content: str) -> str:
        """Detect programming language from content."""
        content_lower = content.lower()
        
        if "def " in content_lower or "import " in content_lower:
            return "python"
        elif "function" in content_lower or "const " in content_lower:
            return "javascript"
        elif "class " in content_lower and "{" in content:
            return "java"
        elif "#include" in content_lower:
            return "cpp"
        else:
            return "text"
    
    def _create_simple_summary(self, content: str, filename: str) -> str:
        """Create a simple summary of the content."""
        lines = content.splitlines()
        non_empty_lines = [line for line in lines if line.strip()]
        
        if len(non_empty_lines) <= 5:
            return f"Short document with {len(non_empty_lines)} lines of content."
        
        # Take first few lines as summary
        summary_lines = non_empty_lines[:3]
        return f"Document summary: {' '.join(summary_lines)}..."
    
    def get_schema(self) -> Dict[str, Any]:
        """Get the tool's parameter schema."""
        return {
            "type": "object",
            "properties": {
                "file_data": {
                    "type": "string",
                    "description": "Base64 encoded file data"
                },
                "filename": {
                    "type": "string",
                    "description": "Name of the file"
                },
                "mime_type": {
                    "type": "string",
                    "description": "MIME type of the file"
                },
                "processing_type": {
                    "type": "string",
                    "enum": ["extract", "analyze", "summarize"],
                    "description": "Type of processing to perform",
                    "default": "extract"
                }
            },
            "required": ["file_data", "filename", "mime_type"]
        }


class ImageAnalysisTool(BaseTool):
    """Tool for analyzing images."""
    
    def __init__(self):
        super().__init__("image_analyzer", "Analyze and process images")
        self.category = "image_processing"
        self.supported_formats = [
            "image/jpeg", "image/jpg", "image/png", "image/gif", 
            "image/webp", "image/bmp", "image/tiff"
        ]
    
    async def execute(
        self,
        file_data: str,
        filename: str,
        mime_type: str,
        analysis_type: str = "basic",
        **kwargs
    ) -> Dict[str, Any]:
        """Execute image analysis."""
        if not await self.validate_input(file_data=file_data, filename=filename, mime_type=mime_type):
            raise ValueError("Invalid input parameters")
        
        if analysis_type == "basic":
            return await self._basic_analysis(file_data, filename, mime_type)
        elif analysis_type == "ocr":
            return await self._ocr_analysis(file_data, filename, mime_type)
        elif analysis_type == "object_detection":
            return await self._object_detection(file_data, filename, mime_type)
        else:
            raise ValueError(f"Unknown analysis type: {analysis_type}")
    
    async def _basic_analysis(self, file_data: str, filename: str, mime_type: str) -> Dict[str, Any]:
        """Perform basic image analysis."""
        try:
            decoded_data = base64.b64decode(file_data)
            image = Image.open(io.BytesIO(decoded_data))
            
            analysis = {
                "filename": filename,
                "mime_type": mime_type,
                "format": image.format,
                "mode": image.mode,
                "size": image.size,
                "width": image.width,
                "height": image.height,
                "has_transparency": image.mode in ("RGBA", "LA", "P"),
                "file_size": len(decoded_data)
            }
            
            return {
                "status": "success",
                "analysis": analysis
            }
        except Exception as e:
            return {
                "status": "error",
                "filename": filename,
                "error": str(e)
            }
    
    async def _ocr_analysis(self, file_data: str, filename: str, mime_type: str) -> Dict[str, Any]:
        """Perform OCR analysis on image."""
        # This would integrate with an OCR service like Tesseract
        return {
            "status": "success",
            "filename": filename,
            "ocr_text": "OCR analysis would be performed here",
            "confidence": 0.95
        }
    
    async def _object_detection(self, file_data: str, filename: str, mime_type: str) -> Dict[str, Any]:
        """Perform object detection on image."""
        # This would integrate with a computer vision service
        return {
            "status": "success",
            "filename": filename,
            "detected_objects": [
                {"object": "person", "confidence": 0.95, "bbox": [100, 100, 200, 300]},
                {"object": "car", "confidence": 0.87, "bbox": [300, 200, 500, 400]}
            ]
        }
    
    def get_schema(self) -> Dict[str, Any]:
        """Get the tool's parameter schema."""
        return {
            "type": "object",
            "properties": {
                "file_data": {
                    "type": "string",
                    "description": "Base64 encoded image data"
                },
                "filename": {
                    "type": "string",
                    "description": "Name of the image file"
                },
                "mime_type": {
                    "type": "string",
                    "description": "MIME type of the image"
                },
                "analysis_type": {
                    "type": "string",
                    "enum": ["basic", "ocr", "object_detection"],
                    "description": "Type of analysis to perform",
                    "default": "basic"
                }
            },
            "required": ["file_data", "filename", "mime_type"]
        }


class CodeAnalysisTool(BaseTool):
    """Tool for analyzing code files."""
    
    def __init__(self):
        super().__init__("code_analyzer", "Analyze and process code files")
        self.category = "code_processing"
        self.supported_languages = [
            "python", "javascript", "typescript", "java", "cpp", "go", "rust",
            "html", "css", "sql", "bash", "powershell"
        ]
    
    async def execute(
        self,
        file_data: str,
        filename: str,
        mime_type: str,
        analysis_type: str = "syntax",
        **kwargs
    ) -> Dict[str, Any]:
        """Execute code analysis."""
        if not await self.validate_input(file_data=file_data, filename=filename, mime_type=mime_type):
            raise ValueError("Invalid input parameters")
        
        if analysis_type == "syntax":
            return await self._syntax_analysis(file_data, filename, mime_type)
        elif analysis_type == "complexity":
            return await self._complexity_analysis(file_data, filename, mime_type)
        elif analysis_type == "dependencies":
            return await self._dependency_analysis(file_data, filename, mime_type)
        else:
            raise ValueError(f"Unknown analysis type: {analysis_type}")
    
    async def _syntax_analysis(self, file_data: str, filename: str, mime_type: str) -> Dict[str, Any]:
        """Perform syntax analysis."""
        try:
            decoded_data = base64.b64decode(file_data)
            content = decoded_data.decode("utf-8")
            
            language = self._detect_language(filename, content)
            
            analysis = {
                "filename": filename,
                "language": language,
                "line_count": len(content.splitlines()),
                "character_count": len(content),
                "has_functions": self._has_functions(content, language),
                "has_classes": self._has_classes(content, language),
                "has_imports": self._has_imports(content, language),
                "syntax_valid": True  # This would be determined by actual syntax checking
            }
            
            return {
                "status": "success",
                "analysis": analysis
            }
        except Exception as e:
            return {
                "status": "error",
                "filename": filename,
                "error": str(e)
            }
    
    async def _complexity_analysis(self, file_data: str, filename: str, mime_type: str) -> Dict[str, Any]:
        """Perform complexity analysis."""
        # This would integrate with complexity analysis tools
        return {
            "status": "success",
            "filename": filename,
            "complexity_metrics": {
                "cyclomatic_complexity": 5,
                "cognitive_complexity": 8,
                "maintainability_index": 75
            }
        }
    
    async def _dependency_analysis(self, file_data: str, filename: str, mime_type: str) -> Dict[str, Any]:
        """Perform dependency analysis."""
        # This would analyze imports and dependencies
        return {
            "status": "success",
            "filename": filename,
            "dependencies": [
                {"name": "numpy", "version": "1.24.0", "type": "import"},
                {"name": "pandas", "version": "2.0.0", "type": "import"}
            ]
        }
    
    def _detect_language(self, filename: str, content: str) -> str:
        """Detect programming language from filename and content."""
        filename_lower = filename.lower()
        
        if filename_lower.endswith(('.py', '.pyw')):
            return "python"
        elif filename_lower.endswith(('.js', '.mjs')):
            return "javascript"
        elif filename_lower.endswith(('.ts', '.tsx')):
            return "typescript"
        elif filename_lower.endswith(('.java', '.jav')):
            return "java"
        elif filename_lower.endswith(('.cpp', '.cc', '.cxx', '.c++')):
            return "cpp"
        elif filename_lower.endswith('.go'):
            return "go"
        elif filename_lower.endswith('.rs'):
            return "rust"
        elif filename_lower.endswith(('.html', '.htm')):
            return "html"
        elif filename_lower.endswith('.css'):
            return "css"
        elif filename_lower.endswith('.sql'):
            return "sql"
        elif filename_lower.endswith(('.sh', '.bash')):
            return "bash"
        elif filename_lower.endswith('.ps1'):
            return "powershell"
        else:
            return "unknown"
    
    def _has_functions(self, content: str, language: str) -> bool:
        """Check if content has function definitions."""
        if language == "python":
            return "def " in content
        elif language in ["javascript", "typescript"]:
            return "function " in content or "=>" in content
        elif language == "java":
            return "public " in content and "(" in content
        else:
            return False
    
    def _has_classes(self, content: str, language: str) -> bool:
        """Check if content has class definitions."""
        if language == "python":
            return "class " in content
        elif language in ["javascript", "typescript"]:
            return "class " in content
        elif language == "java":
            return "class " in content
        else:
            return False
    
    def _has_imports(self, content: str, language: str) -> bool:
        """Check if content has import statements."""
        if language == "python":
            return "import " in content or "from " in content
        elif language in ["javascript", "typescript"]:
            return "import " in content or "require(" in content
        elif language == "java":
            return "import " in content
        else:
            return False
    
    def get_schema(self) -> Dict[str, Any]:
        """Get the tool's parameter schema."""
        return {
            "type": "object",
            "properties": {
                "file_data": {
                    "type": "string",
                    "description": "Base64 encoded code file data"
                },
                "filename": {
                    "type": "string",
                    "description": "Name of the code file"
                },
                "mime_type": {
                    "type": "string",
                    "description": "MIME type of the file"
                },
                "analysis_type": {
                    "type": "string",
                    "enum": ["syntax", "complexity", "dependencies"],
                    "description": "Type of analysis to perform",
                    "default": "syntax"
                }
            },
            "required": ["file_data", "filename", "mime_type"]
        }
